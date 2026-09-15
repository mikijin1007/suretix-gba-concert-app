from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/charliewu/Documents/香港票务酒店套餐平台')
SOURCE = ROOT / 'PRD_大湾区观演一站式服务平台_v0.2.md'
OUTPUT = ROOT / 'PRD_大湾区观演一站式服务平台_v0.2.docx'

FONT_LATIN = 'Calibri'
FONT_CJK = 'Hiragino Sans GB'
BLUE = '2E74B5'
DARK_BLUE = '1F4D78'
INK = '1F2937'
MUTED = '667085'
LIGHT_BLUE = 'E8EEF5'
LIGHT_GRAY = 'F4F6F9'
BORDER = 'CBD5E1'
RISK_RED = '9B1C1C'
CAUTION_GOLD = '7A5A00'
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    name = 'Menlo' if mono else FONT_LATIN
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    fonts.set(qn('w:ascii'), name)
    fonts.set(qn('w:hAnsi'), name)
    fonts.set(qn('w:eastAsia'), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_table_borders(table, color=BORDER, size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = borders.find(qn(f'w:{edge}'))
        if tag is None:
            tag = OxmlElement(f'w:{edge}')
            borders.append(tag)
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), size)
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), color)


def set_table_geometry(table, widths):
    assert sum(widths) == TABLE_WIDTH, widths
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(TABLE_WIDTH))
    tbl_w.set(qn('w:type'), 'dxa')

    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), str(TABLE_INDENT))
    tbl_ind.set(qn('w:type'), 'dxa')

    layout = tbl_pr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tbl_pr.append(layout)
    layout.set(qn('w:type'), 'fixed')

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(widths[idx]))
            tc_w.set(qn('w:type'), 'dxa')
            cell.width = Inches(widths[idx] / 1440)


def table_widths(column_count):
    patterns = {
        2: [2100, 7260],
        3: [1600, 3100, 4660],
        4: [1100, 1900, 3060, 3300],
        5: [800, 1600, 2200, 2260, 2500],
    }
    if column_count in patterns:
        return patterns[column_count]
    base = TABLE_WIDTH // column_count
    widths = [base] * column_count
    widths[-1] += TABLE_WIDTH - sum(widths)
    return widths


def paragraph_border_bottom(paragraph, color=BLUE, size='12', space='6'):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is None:
        p_bdr = OxmlElement('w:pBdr')
        p_pr.append(p_bdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color)
    p_bdr.append(bottom)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('第 ')
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    end = paragraph.add_run(' 页')
    set_run_font(end, size=9, color=MUTED)


def define_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))]
    existing_num = [int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def make_abstract(num_id, fmt, text, bullet_font=None):
        abstract = OxmlElement('w:abstractNum')
        abstract.set(qn('w:abstractNumId'), str(num_id))
        multi = OxmlElement('w:multiLevelType')
        multi.set(qn('w:val'), 'singleLevel')
        abstract.append(multi)
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), '0')
        start = OxmlElement('w:start')
        start.set(qn('w:val'), '1')
        lvl.append(start)
        num_fmt = OxmlElement('w:numFmt')
        num_fmt.set(qn('w:val'), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement('w:lvlText')
        lvl_text.set(qn('w:val'), text)
        lvl.append(lvl_text)
        jc = OxmlElement('w:lvlJc')
        jc.set(qn('w:val'), 'left')
        lvl.append(jc)
        p_pr = OxmlElement('w:pPr')
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'num')
        tab.set(qn('w:pos'), '540')
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '540')
        ind.set(qn('w:hanging'), '270')
        p_pr.append(ind)
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:after'), '80')
        spacing.set(qn('w:line'), '300')
        spacing.set(qn('w:lineRule'), 'auto')
        p_pr.append(spacing)
        lvl.append(p_pr)
        if bullet_font:
            r_pr = OxmlElement('w:rPr')
            fonts = OxmlElement('w:rFonts')
            fonts.set(qn('w:ascii'), bullet_font)
            fonts.set(qn('w:hAnsi'), bullet_font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    make_abstract(next_abs, 'bullet', '•', 'Symbol')
    bullet_abs = next_abs
    next_abs += 1
    make_abstract(next_abs, 'decimal', '%1.')
    decimal_abs = next_abs

    def make_num(abstract_id):
        nonlocal next_num
        num = OxmlElement('w:num')
        num.set(qn('w:numId'), str(next_num))
        abstract_id_node = OxmlElement('w:abstractNumId')
        abstract_id_node.set(qn('w:val'), str(abstract_id))
        num.append(abstract_id_node)
        # LibreOffice may otherwise continue numbering across distinct list
        # instances that share one abstract definition. Force every instance
        # to restart at 1 so each Markdown ordered-list block is independent.
        lvl_override = OxmlElement('w:lvlOverride')
        lvl_override.set(qn('w:ilvl'), '0')
        start_override = OxmlElement('w:startOverride')
        start_override.set(qn('w:val'), '1')
        lvl_override.append(start_override)
        num.append(lvl_override)
        numbering.append(num)
        current = next_num
        next_num += 1
        return current

    bullet_num = make_num(bullet_abs)

    def new_decimal_num():
        return make_num(decimal_abs)

    return bullet_num, new_decimal_num


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn('w:numPr'))
    if num_pr is None:
        num_pr = OxmlElement('w:numPr')
        p_pr.append(num_pr)
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    num_id_el = OxmlElement('w:numId')
    num_id_el.set(qn('w:val'), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def add_inline(paragraph, text, *, base_size=11, base_color=INK):
    token_re = re.compile(r'(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))')
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=base_color)
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size - 0.5, color=DARK_BLUE, mono=True)
            run.font.highlight_color = None
        else:
            link_match = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            label, url = link_match.groups()
            run = paragraph.add_run(f'{label}（{url}）')
            set_run_font(run, size=base_size, color=BLUE)
            run.underline = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, color=base_color)


def add_markdown_table(doc, rows):
    headers = rows[0]
    body = rows[2:] if len(rows) > 1 and all(re.fullmatch(r':?-{3,}:?', c.strip()) for c in rows[1]) else rows[1:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, table_widths(len(headers)))
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        add_inline(p, header.strip(), base_size=9.5, base_color=DARK_BLUE)
        for r in p.runs:
            r.bold = True

    for row_data in body:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, value.strip(), base_size=9.2, base_color=INK)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def parse_table_row(line):
    stripped = line.strip().strip('|')
    return [part.strip() for part in stripped.split('|')]


def configure_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn('w:ascii'), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ('Heading 1', 16, BLUE, 18, 10),
        ('Heading 2', 13, BLUE, 14, 7),
        ('Heading 3', 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn('w:ascii'), FONT_LATIN)
        style._element.rPr.rFonts.set(qn('w:hAnsi'), FONT_LATIN)
        style._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0


def set_page_layout(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = hp.add_run('ZC Digitals  |  大湾区观演一站式服务平台')
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = hp.add_run('    PRD v0.2 · 内部评审稿')
    set_run_font(right, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    add_page_number(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('产品需求文档')
    set_run_font(r, size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('大湾区观演一站式服务平台')
    set_run_font(r, size=26, bold=True, color='111827')

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run('PRD v0.2｜管理型二级票务平台＋票酒套餐')
    set_run_font(r, size=14, color=MUTED)

    metadata = [
        ('编制', 'ZC Digitals（桢诚数科）'),
        ('日期', '2026-08-17'),
        ('状态', '内部评审稿'),
        ('密级', '内部资料，未经授权不得外传'),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        rl = p.add_run(f'{label}：')
        set_run_font(rl, size=10.5, bold=True, color='111827')
        rv = p.add_run(value)
        set_run_font(rv, size=10.5, color='374151')

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(12)
    rule.paragraph_format.space_after = Pt(16)
    paragraph_border_bottom(rule, color=BLUE, size='16', space='6')

    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [TABLE_WIDTH])
    set_table_borders(callout, color='D6E4F0', size='6')
    cell = callout.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('本版关键结论：')
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    add_inline(
        p,
        '平台统一收款并对用户负责；一期强制“票＋酒店”；票源分为现票、承诺库存和询价库存；供应商采用邀请制、分级额度和滚动风险金；现场无法入场执行门票与未使用服务退款，并按500元/票、2,000元/单封顶赔付。',
        base_size=10.5,
        base_color=INK,
    )
    doc.add_page_break()


def add_contents(doc):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(0)
    p.add_run('目录')
    entries = [
        ('1', '执行摘要'), ('2', '用户、角色与商业关系'), ('3', '产品范围与优先级'),
        ('4', '核心交易与库存规则'), ('5', '酒店与票酒套餐规则'), ('6', '供应商准入、信用与结算'),
        ('7', '用户担保、补票、退款与赔付'), ('8', 'C端功能需求'), ('9', '供应商工作台需求'),
        ('10', '运营后台需求'), ('11', '状态机'), ('12', '数据模型'), ('13', '支付、财务与对账'),
        ('14', '安全、隐私与合规基线'), ('15', '非功能需求'), ('16', '数据指标与埋点'),
        ('17', '一期验收标准'), ('18', '上线运营机制'), ('19', '待决策项（TBD）'),
        ('附录A', '关键术语'), ('附录B', '退款与赔付矩阵'), ('附录C', '原型对齐清单'),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1600, 7760])
    set_table_borders(table, color='E5E7EB', size='4')
    table.rows[0]._element.getparent().remove(table.rows[0]._element)
    for number, title in entries:
        row = table.add_row()
        for idx, value in enumerate((number, title)):
            cell = row.cells[idx]
            set_cell_margins(cell, top=55, bottom=55, start=100, end=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=9.5, bold=(idx == 0), color=DARK_BLUE if idx == 0 else INK)
    doc.add_page_break()


def render_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    in_code = False
    code_lines = []
    bullet_num, new_decimal_num = define_numbering(doc)
    current_decimal_num = None
    in_decimal_block = False
    seen_first_h1 = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                for code_line in code_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.18)
                    p.paragraph_format.right_indent = Inches(0.18)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.05
                    p_pr = p._p.get_or_add_pPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), 'F3F4F6')
                    p_pr.append(shd)
                    r = p.add_run(code_line or ' ')
                    set_run_font(r, size=9, color='374151', mono=True)
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped or stripped == '---':
            in_decimal_block = False
            i += 1
            continue

        if stripped.startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(parse_table_row(lines[i]))
                i += 1
            add_markdown_table(doc, table_lines)
            in_decimal_block = False
            continue

        heading = re.match(r'^(#{1,3})\s+(.+)$', stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1 and title == '大湾区观演一站式服务平台':
                i += 1
                continue
            if level == 2 and title == '产品需求文档（PRD）v0.2':
                i += 1
                continue
            style = 'Heading 1' if level == 1 else ('Heading 2' if level == 2 else 'Heading 3')
            if level == 1:
                if seen_first_h1:
                    doc.add_page_break()
                seen_first_h1 = True
            p = doc.add_paragraph(style=style)
            add_inline(p, title, base_size=16 if level == 1 else (13 if level == 2 else 12), base_color=BLUE if level < 3 else DARK_BLUE)
            for r in p.runs:
                r.bold = True
            in_decimal_block = False
            i += 1
            continue

        if re.match(r'^\*\*[^*]+：\*\*', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, stripped)
            i += 1
            continue

        if stripped.startswith('> '):
            table = doc.add_table(rows=1, cols=1)
            set_table_geometry(table, [TABLE_WIDTH])
            set_table_borders(table, color='BFD4E8', size='6')
            cell = table.cell(0, 0)
            set_cell_shading(cell, LIGHT_GRAY)
            set_cell_margins(cell, top=120, bottom=120, start=170, end=170)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, stripped[2:], base_size=10.5, base_color=DARK_BLUE)
            for r in p.runs:
                r.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if stripped.startswith('- '):
            p = doc.add_paragraph()
            apply_num(p, bullet_num)
            add_inline(p, stripped[2:])
            in_decimal_block = False
            i += 1
            continue

        ordered = re.match(r'^\d+\.\s+(.+)$', stripped)
        if ordered:
            if not in_decimal_block:
                current_decimal_num = new_decimal_num()
                in_decimal_block = True
            p = doc.add_paragraph()
            apply_num(p, current_decimal_num)
            add_inline(p, ordered.group(1))
            i += 1
            continue

        if stripped == '**文档结束**':
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        in_decimal_block = False
        i += 1


def build():
    markdown = SOURCE.read_text(encoding='utf-8')
    # Skip front-matter content already represented by the Word cover, then resume at change log.
    start = markdown.index('## 变更记录')
    body = markdown[start:]

    doc = Document()
    configure_styles(doc)
    set_page_layout(doc)
    core = doc.core_properties
    core.title = '大湾区观演一站式服务平台 · 产品需求文档（PRD）v0.2'
    core.subject = '管理型二级票务平台与票酒套餐产品需求'
    core.author = 'ZC Digitals（桢诚数科）'
    core.keywords = 'PRD, 票务, 酒店, 香港演出, 供应商, 担保履约'
    core.comments = '内部评审稿'

    add_cover(doc)
    add_contents(doc)
    render_markdown(doc, body)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
